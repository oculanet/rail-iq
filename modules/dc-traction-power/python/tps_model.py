"""
DC Traction Power Simulation model - command line program
==========================================================
Usage examples (run from this folder):

  python tps_model.py --make-templates
        writes Alignment_Template.xlsx and Train_Data_Template.docx to fill in

  python tps_model.py --alignment Alignment.xlsx --train Train.docx --out results
        reads the alignment and train data, simulates normal service, writes results

  python tps_model.py --inputs my_inputs.json --design-check --out results
        inputs exported from the HTML model ("Export inputs"), plus all single-TSS outages

  python tps_model.py --n-up 2 --n-down 2 --tss-count 9 --tss-first 300 --tss-spacing 2300 --out results
        regenerate tracks and traction substations from counts / first chainage / spacing

  python tps_model.py --add-tss "TSS-X,9800,3000,3300" --outage TSS4 --out results

Every input can also be edited in the inputs JSON file (save one with --save-inputs).
The same JSON drives the HTML model, and both give the same results.
"""
from __future__ import annotations

import argparse
import json
import math
import os
import sys
import time

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import tps_engine as E  # noqa: E402
import tps_readers as R  # noqa: E402
import tps_reports as RP  # noqa: E402


# ----------------------------------------------------------------------
# inputs
# ----------------------------------------------------------------------
def load_inputs(path):
    if not path:
        return E.defaults()
    with open(path, encoding="utf-8") as f:
        cfg = json.load(f)
    base = E.defaults()
    for k, v in base.items():          # fill any section missing from an older file
        cfg.setdefault(k, v)
    return cfg


def apply_cli(cfg, a, log):
    if a.alignment:
        part, lg = R.read_alignment_xlsx(a.alignment)
        R.merge_alignment(cfg, part)
        cfg["alignment"]["source"] = os.path.basename(a.alignment)
        log += [f"[alignment] {x}" for x in lg]
    if a.train:
        i = a.train_index
        while len(cfg["trains"]) <= i:
            cfg["trains"].append(json.loads(json.dumps(cfg["trains"][-1])))
        axes = tuple(float(x) for x in a.graph_axes.split(",")) if a.graph_axes else None
        T, rep = R.read_train_docx(a.train, base=cfg["trains"][i], graph_axes=axes, graph_kind=a.graph_use, overlay_dir=a.out)
        if not T.get("name") or T.get("name") == cfg["trains"][i].get("name"):
            T["name"] = os.path.splitext(os.path.basename(a.train))[0]
        cfg["trains"][i] = T
        for p in rep["params"]:
            log.append(f"[train] {p[0]} = {p[1]}    <- \"{p[2]}\"")
        for c in rep["curves"]:
            log.append(f"[train] curve {c[0]}: {c[1]} points from {c[3]} ({c[2]})")
        log += [f"[train] {x}" for x in rep["log"]]
    if a.n_up is not None or a.n_down is not None:
        if a.n_up is not None:
            cfg["track_gen"]["n_up"] = a.n_up
        if a.n_down is not None:
            cfg["track_gen"]["n_down"] = a.n_down
        E.generate_tracks(cfg)
    if a.headway:
        for t in cfg["tracks"]:
            t["headway"] = a.headway
    if a.tss_count is not None or a.tss_first is not None or a.tss_spacing is not None:
        g = cfg["tss_gen"]
        if a.tss_count is not None:
            g["count"] = a.tss_count
        if a.tss_first is not None:
            g["first_ch"] = a.tss_first
        if a.tss_spacing is not None:
            g["spacing"] = a.tss_spacing
        E.generate_tss(cfg)
    for s in a.add_tss or []:
        p = [x.strip() for x in s.split(",")]
        cfg["tss"].append({"name": p[0], "ch": float(p[1]), "rect_kW": float(p[2]) if len(p) > 2 else 3000,
                           "xfmr_kVA": float(p[3]) if len(p) > 3 else 3300, "V0": "", "Req_mohm": "", "inv_kW": 0, "in_service": True})
        cfg["tss"].sort(key=lambda t: float(t["ch"]))
    if a.period:
        cfg["sim"]["period"] = a.period
    return cfg


# ----------------------------------------------------------------------
# outputs
# ----------------------------------------------------------------------
def write_excel(path, cfg, r, flt, dc=None, st=None, timeseries=True):
    import xlsxwriter
    wb = xlsxwriter.Workbook(path, {"nan_inf_to_errors": True})
    hd = wb.add_format({"bold": True, "bg_color": "#1F3A4D", "font_color": "white", "border": 1, "text_wrap": True, "valign": "top"})
    n0 = wb.add_format({"num_format": "#,##0", "border": 1})
    n1 = wb.add_format({"num_format": "#,##0.0", "border": 1})
    n3 = wb.add_format({"num_format": "0.000", "border": 1})
    tx = wb.add_format({"border": 1})
    ok = wb.add_format({"bold": True, "font_color": "#1B7F3B", "border": 1})
    bad = wb.add_format({"bold": True, "font_color": "#B3261E", "bg_color": "#FBE3E1", "border": 1})
    big = wb.add_format({"bold": True, "font_size": 14})

    def table(ws, row, heads, data, fmts, widths=None):
        for j, h in enumerate(heads):
            ws.write(row, j, h, hd)
        for i, d in enumerate(data):
            for j, v in enumerate(d):
                f = fmts[j] if j < len(fmts) else tx
                if isinstance(v, bool) or v in ("PASS", "FAIL"):
                    ws.write(row + 1 + i, j, v if not isinstance(v, bool) else ("PASS" if v else "FAIL"),
                             ok if (v is True or v == "PASS") else bad)
                elif v is None or (isinstance(v, float) and not math.isfinite(v)):
                    ws.write_blank(row + 1 + i, j, None, tx)
                else:
                    ws.write(row + 1 + i, j, v, f)
        if widths:
            for j, w in enumerate(widths):
                ws.set_column(j, j, w)
        return row + len(data) + 2

    ch = r["checks"]
    ws = wb.add_worksheet("Summary")
    ws.write(0, 0, cfg["project"]["name"], big)
    ws.write(1, 0, f"{cfg['project'].get('system', '')}   |   scenario: {r['scenario']}   |   {r['engine']}")
    fails = [c for c in ch if not c["pass"]]
    ws.write(2, 0, f"Checks: {len(ch) - len(fails)} pass, {len(fails)} fail", bad if fails else ok)
    kv = [
        ("Simulated period", r["period"], "s"), ("Minimum train voltage", r["trainV"]["min"], "V"),
        ("Maximum train voltage", r["trainV"]["max"], "V"), ("Mean useful voltage", r["trainV"]["mean_useful"], "V"),
        ("Energy from TSS (DC)", r["energy"]["tss_kWh"], "kWh"), ("Regenerated energy returned to line", r["energy"]["regen_out_kWh"], "kWh"),
        ("Energy to rheostats (not accepted)", r["energy"]["rheostat_kWh"], "kWh"), ("Inverter energy returned", r["energy"]["inverter_kWh"], "kWh"),
        ("Conductor + return + feeder losses", r["energy"]["losses_kWh"], "kWh"),
        ("Max fault current", max(t["max"] for t in flt["tracks"]) / 1000, "kA"),
        ("Min fault current (worst feeder)", min(t["minFeeder"] for t in flt["tracks"]) / 1000, "kA"),
        ("Solver: max iterations / non-converged steps", f"{r['solver']['iter_max_used']} / {r['solver']['non_converged_steps']}", ""),
    ]
    table(ws, 4, ["Quantity", "Value", "Unit"], kv, [tx, n1, tx], [44, 16, 8])
    row = 4 + len(kv) + 3
    ws.write(row - 1, 0, "Tracks and service", big)
    row = table(ws, row, ["Track", "Direction", "Headway (s)", "Trains / h", "Trip time (min)", "Trains on track"],
                [[t["name"], t["dir"], t["headway"], t["tph"], t["trip"] / 60, t["inService"]] for t in r["tracks"]],
                [tx, tx, n0, n1, n1, n0])
    for w in r["warnings"]:
        ws.write(row, 0, "Warning: " + w, bad)
        row += 1

    ws = wb.add_worksheet("Checks")
    table(ws, 0, ["Result", "Group", "Check", "Location", "Value", "Limit", "Unit", "Utilisation %"],
          [["PASS" if c["pass"] else "FAIL", c["group"], c["item"], c["where"], c["value"], c["limit"], c["unit"], c["ratio"] * 100] for c in ch],
          [tx, tx, tx, tx, n1, n1, tx, n1], [8, 16, 52, 26, 12, 12, 6, 13])

    ws = wb.add_worksheet("TSS")
    d = r["contingency"] and cfg["limits"]["contingency"] or cfg["limits"]["normal"]
    Vn = E.num(cfg["electrical"].get("V_nom"), 750)

    def std_list(txt, default):
        try:
            v = sorted(float(x) for x in str(txt).replace(";", ",").split(",") if x.strip())
            return v or default
        except ValueError:
            return default
    sr = std_list(cfg["limits"].get("std_rect_kW", ""), [1000, 1500, 2000, 2500, 3000, 3300, 3500, 4000, 4500, 5000, 6000, 7000, 8000])
    sx = std_list(cfg["limits"].get("std_xfmr_kVA", ""), [1250, 1600, 2000, 2500, 3150, 3300, 4000, 4500, 5000, 6300, 8000])
    nxt = lambda v, lst: next((x for x in lst if x >= v), None)
    rows = []
    for t in r["tss"]:
        rows.append([t["name"], t["ch"], "in service" if t["active"] else "OUT", t["rect_kW"], t["Ir"], t["xfmr_kVA"], t["V0"], t["Req"] * 1000,
                     t["Ipk"], t["I1"], t["I15"], t["I1h"], t["Iavg"],
                     t["I1h"] / t["Ir"] * 100 if t["active"] else None, t["I15"] / t["Ir"] * 100 if t["active"] else None,
                     t["S1h"], t["S15"], t["S1h"] / max(t["xfmr_kVA"], 1) * 100 if t["active"] else None,
                     t["Pavg"], t["Ppk"], t["E_kWh"], t["Einv_kWh"], t["Vmin"], t["Vmax"]])
        if t["active"]:
            need_kW = max(t["I1h"] / d["h1"], t["I15"] / d["m15"], t["I1"] / d["m1"]) * Vn / 1000
            need_kVA = max(t["S1h"] / d["h1"], t["S15"] / d["m15"], t["S1"] / d["m1"])
            rows[-1] += [need_kW, nxt(need_kW, sr), need_kVA, nxt(need_kVA, sx)]
    table(ws, 0, ["TSS", "Chainage (m)", "Status", "Rectifier (kW)", "Rated I (A)", "Transformer (kVA)", "No-load V", "Source R (mΩ)",
                  "Peak I (A)", "1-min RMS (A)", "15-min RMS (A)", "1-h RMS (A)", "Mean I (A)", "1-h % of rating", "15-min % of rating",
                  "Transformer 1-h (kVA)", "Transformer 15-min (kVA)", "Transformer 1-h %", "Mean P (kW)", "Peak P (kW)", "Energy (kWh)",
                  "Inverter energy (kWh)", "Min bus V", "Max bus V", "Rectifier needed (kW)", "Next standard (kW)",
                  "Transformer needed (kVA)", "Next standard (kVA)"],
          rows, [tx, n0, tx, n0, n0, n0, n0, n3] + [n0] * 5 + [n1, n1, n0, n0, n1, n0, n0, n0, n0, n1, n1, n0, n0, n0, n0], [10] + [12] * 27)
    ws.write(len(rows) + 2, 0, f"Duty multipliers used: 1-h {d['h1']}, 15-min {d['m15']}, 1-min {d['m1']} × rating")

    ws = wb.add_worksheet("Feeders & rail")
    rr = table(ws, 0, ["Feeder: TSS", "Track", "Peak I (A)", "1-h RMS (A)", "15-min RMS (A)"],
               [[f["tss"], f["track"], f["Ipk"], f["I1h"], f["I15"]] for f in r["feeders"]], [tx, tx, n0, n0, n0], [14, 10, 12, 12, 14])
    table(ws, rr + 1, ["Conductor rail at TSS", "Track", "Side", "Peak I (A)", "1-h RMS (A)", "15-min RMS (A)"],
          [[c["tss"], c["track"], c["side"], c["Ipk"], c["I1h"], c["I15"]] for c in r["conductor"]], [tx, tx, tx, n0, n0, n0])

    ws = wb.add_worksheet("Train runs")
    row = 0
    for key, run in r["runs"].items():
        ws.write(row, 0, f"{run['train']}  {run['dir']}  trip {run['T'] / 60:.2f} min  traction {run['E_trac_kWh']:.1f} kWh  regen {run['E_regen_kWh']:.1f} kWh  aux {run['E_aux_kWh']:.1f} kWh", big)
        row = table(ws, row + 1, ["Station", "Chainage (m)", "Arrive (s)", "Depart (s)", "Arrive (min)"],
                    [[s["name"], s["ch"], s["arr"], s["dep"], s["arr"] / 60] for s in run["stations"]], [tx, n0, n0, n0, n1], [14, 12, 12, 12, 12])
        if run["warn"]:
            ws.write(row - 1, 0, run["warn"], bad)

    ws = wb.add_worksheet("TSS current time series")
    heads = ["Time (s)"] + [f"{t['name']} I (A)" for t in r["tss"]] + [f"{t['name']} V" for t in r["tss"]] + ["Min train V", "Max train V", "Trains"]
    for j, h in enumerate(heads):
        ws.write(0, j, h, hd)
    tsI, tsV = r["ts"]["tssI"], r["ts"]["tssV"]
    nS = tsI.shape[1] if timeseries else 1
    nT = len(r["tss"])
    for s in range(nS):
        ws.write(s + 1, 0, s * r["dt"])
        for i in range(nT):
            ws.write(s + 1, 1 + i, round(float(tsI[i, s]), 1))
            v = float(tsV[i, s])
            if math.isfinite(v):
                ws.write(s + 1, 1 + nT + i, round(v, 1))
        a, b = float(r["ts"]["vTrainMin"][s]), float(r["ts"]["vTrainMax"][s])
        if math.isfinite(a):
            ws.write(s + 1, 1 + 2 * nT, round(a, 1))
            ws.write(s + 1, 2 + 2 * nT, round(b, 1))
        ws.write(s + 1, 3 + 2 * nT, int(r["ts"]["nTrains"][s]))

    ws = wb.add_worksheet("Voltage envelope")
    col = 0
    for e in r["trainV"]["envelope"]:
        for j, h in enumerate([f"{e['track']} chainage (m)", "Min V", "Mean V", "Max V"]):
            ws.write(0, col + j, h, hd)
        for i in range(len(e["x"])):
            ws.write_row(i + 1, col, [e["x"][i], round(e["min"][i], 1), round(e["mean"][i], 1), round(e["max"][i], 1)])
        col += 5

    ws = wb.add_worksheet("Short circuit")
    col = 0
    for t in flt["tracks"]:
        for j, h in enumerate([f"{t['track']} fault chainage (m)", "Total fault current (A)", "Largest feeder contribution (A)"]):
            ws.write(0, col + j, h, hd)
        for i in range(len(t["x"])):
            ws.write_row(i + 1, col, [t["x"][i], round(t["If"][i]), round(t["Ifeeder"][i])])
        col += 4

    if dc:
        ws = wb.add_worksheet("Design check")
        table(ws, 0, ["Scenario", "Min train V", "Max train V", "Worst rectifier %", "Worst transformer %", "Worst conductor %",
                      "Worst feeder %", "Min feeder fault (A)", "Energy (kWh)", "Failures", "First failure"],
              [[x["scenario"], x["vmin"], x["vmax"], x["rect"] * 100, x["xfmr"] * 100, x["cond"] * 100, x["feeder"] * 100,
                x["fault_min"], x["E_kWh"], x["nFail"], x["firstFail"]] for x in dc],
              [tx, n0, n0, n1, n1, n1, n1, n0, n0, n0, tx], [22, 11, 11, 13, 13, 13, 12, 13, 12, 9, 60])
        rows = []
        for x in dc:
            for c in x["checks"]:
                if not c["pass"]:
                    rows.append([x["scenario"], c["group"], c["item"], c["where"], c["value"], c["limit"], c["unit"], c["ratio"] * 100])
        ws = wb.add_worksheet("Design check failures")
        table(ws, 0, ["Scenario", "Group", "Check", "Location", "Value", "Limit", "Unit", "Utilisation %"], rows,
              [tx, tx, tx, tx, n1, n1, tx, n1], [22, 16, 50, 26, 12, 12, 6, 13])

    if st:
        add_study_sheets(wb, table, hd, flt, st, tx, n0, n1, n3)
    ws = wb.add_worksheet("Inputs")
    ws.set_column(0, 0, 120)
    for i, line in enumerate(json.dumps(cfg, indent=1).splitlines()):
        ws.write(i, 0, line)
    wb.close()


def add_study_sheets(wb, table, hd, flt, st, tx, n0, n1, n3):
    rows = []
    for t in flt["tracks"]:
        for i, x in enumerate(t["x"]):
            rows.append([t["track"], x, t["If"][i], t["Ifeeder"][i], t["didt"][i], t["didtFeeder"][i], t["tau_ms"][i], t["ttrip_ms"][i], t["Urail"][i]])
    ws = wb.add_worksheet("SC along line")
    table(ws, 0, ["Track", "Chainage (m)", "Total fault (A)", "Largest feeder (A)", "di/dt total (A/ms)", "di/dt feeder (A/ms)", "Time constant (ms)",
                  "Time to trip (ms)", "Rail potential (V)"], rows, [tx, n0, n0, n0, n0, n0, n3, n1, n0], [8, 12, 14, 16, 16, 16, 16, 14, 14])
    ws = wb.add_worksheet("SC per TSS")
    table(ws, 0, ["TSS", "HV Ik'' (kA)", "HV ip (kA)", "Valve Ik'' (kA)", "Valve ip (kA)", "Xc (mOhm)", "Source R (mOhm)", "R from AC (mOhm)",
                  "Source L (mH)", "DC terminal fault (kA)", "di/dt (A/ms)", "tau (ms)"],
          [[a["name"], a["IkHV_kA"], a["ipHV_kA"], a["IkLV_kA"], a["ipLV_kA"], a["Xc_mohm"], a["Req_mohm"], a["ReqAC_mohm"], a["Ldc_mH"], a["IkD_kA"],
            a["didtD_A_ms"], a["tauD_ms"]] for a in flt["tssAC"]], [tx, n3, n3, n1, n1, n3, n3, n3, n3, n1, n0, n3], [10] + [14] * 11)
    H = st["harm"]
    ents = ([H["common"]] if H["common"] else []) + [t for t in H["tss"] if t["active"]]
    ws = wb.add_worksheet("Harmonics")
    rr = table(ws, 0, ["Supply point", "Voltage THD 95% (%)", "Voltage THD max (%)", "Current THD rated (%)", "Overlap rated (deg)"],
               [[t["name"], t["THDv95"], t["THDvmax"], t.get("THDi_rated"), t.get("mu_rated_deg")] for t in ents], [tx, n3, n3, n1, n1], [18, 16, 16, 18, 16])
    table(ws, rr + 1, ["Order h", "Limit (%)"] + [t["name"] + " Vh95 (%)" for t in ents],
          [[h, H["limits"][j]] + [t["Vh95"][j] for t in ents] for j, h in enumerate(H["orders"])], [n0, n3] + [n3] * len(ents))
    em = st["emc"]
    ws = wb.add_worksheet("EMC rail potential")
    table(ws, 0, ["Chainage (m)", "Highest (V)", "Lowest (V)", "Highest 300 s avg |U| (V)", "Mean |U| (V)"],
          [[x, em["Umax"][j], em["Umin"][j], em["U300"][j], em["Umean"][j]] for j, x in enumerate(em["x"])], [n0, n1, n1, n1, n1], [12, 12, 12, 22, 12])
    ws = wb.add_worksheet("EMC field & telecom")
    rr = table(ws, 0, ["Distance (m)", "B peak (uT)", "B 1-h RMS (uT)"], [[f["d"], f["Bpk_uT"], f["Brms_uT"]] for f in em["field"]], [n1, n1, n1], [16, 14, 14])
    rr = table(ws, rr + 1, ["Section", "Length (m)", "f (Hz)", "Ripple (A)", "Longitudinal emf (mV)", "Weight", "Noise (mV)"],
               [[t["tss"] if k == 0 else "", t["Lpar"] if k == 0 else None, e["f"], e["I"], e["E_mV"], e["w"], t["ep_mV"] if k == 0 else None]
                for t in em["telecom"] for k, e in enumerate(t["E"])], [tx, n0, n0, n3, n1, n3, n3])
    table(ws, rr + 1, ["Quantity", "Value"], [["Stray current, average (A)", em["strayAvg_A"]], ["Stray current (mA/m)", em["stray_mA_m"]],
                                              ["Mutual inductance (uH/m)", em["Mprime_uH_m"]]], [tx, n3])


def write_charts(outdir, cfg, r, flt):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 9, "axes.grid": True, "grid.alpha": 0.3, "axes.spines.top": False, "axes.spines.right": False})
    files = []
    pal = ["#2b6cb0", "#c05621", "#2f855a", "#805ad5", "#b83280", "#2c7a7b", "#975a16", "#4a5568"]

    # long section + speed
    el = E.elevation_profile(cfg, 20)
    fig, ax = plt.subplots(3, 1, figsize=(12, 7.5), sharex=True, gridspec_kw={"height_ratios": [2, 1, 2]})
    ax[0].plot(el["x"], el["z"], color="#4a5568")
    ax[0].set_ylabel("Level (m)")
    for s in cfg["alignment"]["stations"]:
        ax[0].axvline(s["ch"], color="#a0aec0", lw=0.6)
        ax[0].text(s["ch"], max(el["z"]), s["name"], rotation=90, fontsize=7, va="top", ha="right")
    for t in cfg["tss"]:
        ax[0].plot(t["ch"], min(el["z"]), marker="s", color="#c05621")
    ax[0].set_title("Long section, curves and speed profile")
    for c in cfg["alignment"]["curves"]:
        ax[1].fill_between([c["from"], c["to"]], 0, 1000 / max(c["R"], 1), color="#805ad5", alpha=0.5)
        ax[1].text((c["from"] + c["to"]) / 2, 1000 / max(c["R"], 1), f"R{c['R']:.0f}", fontsize=6, ha="center", va="bottom")
    ax[1].set_ylabel("Curvature 1000/R")
    for j, (k, run) in enumerate(r["runs"].items()):
        ax[2].plot(run["x"], run["v"], color=pal[j % len(pal)], lw=1, label=f"{run['train']} {run['dir']}")
    ax[2].set_ylabel("Speed (km/h)")
    ax[2].set_xlabel("Chainage (m)")
    ax[2].legend(loc="lower center", fontsize=7)
    fig.tight_layout()
    p = os.path.join(outdir, "01_long_section_speed.png")
    fig.savefig(p, dpi=130)
    plt.close(fig)
    files.append(p)

    # TSS currents
    fig, ax = plt.subplots(figsize=(12, 4.5))
    t = np.arange(r["ts"]["tssI"].shape[1]) * r["dt"] / 60
    for i, ts in enumerate(r["tss"]):
        if ts["active"]:
            ax.plot(t, r["ts"]["tssI"][i], lw=0.7, color=pal[i % len(pal)], label=ts["name"])
    ax.set_xlabel("Time (min)")
    ax.set_ylabel("Rectifier current (A)")
    ax.set_title(f"TSS currents - {r['scenario']}")
    ax.legend(ncol=8, fontsize=7)
    fig.tight_layout()
    p = os.path.join(outdir, "02_tss_currents.png")
    fig.savefig(p, dpi=130)
    plt.close(fig)
    files.append(p)

    # TSS loading bars
    fig, ax = plt.subplots(figsize=(12, 4))
    names = [x["name"] for x in r["tss"]]
    xx = np.arange(len(names))
    ax.bar(xx - 0.27, [x["I1h"] / x["Ir"] * 100 if x["active"] else 0 for x in r["tss"]], 0.27, label="1-h RMS % rating", color="#2b6cb0")
    ax.bar(xx, [x["I15"] / x["Ir"] * 100 if x["active"] else 0 for x in r["tss"]], 0.27, label="15-min RMS %", color="#63b3ed")
    ax.bar(xx + 0.27, [x["S1h"] / max(x["xfmr_kVA"], 1) * 100 if x["active"] else 0 for x in r["tss"]], 0.27, label="Transformer 1-h %", color="#c05621")
    d = cfg["limits"]["contingency"] if r["contingency"] else cfg["limits"]["normal"]
    ax.axhline(d["h1"] * 100, color="#b3261e", ls="--", lw=1, label="1-h limit")
    ax.set_xticks(xx, names)
    ax.set_ylabel("% of rating")
    ax.legend(fontsize=7)
    ax.set_title("Rectifier and transformer loading")
    fig.tight_layout()
    p = os.path.join(outdir, "03_tss_loading.png")
    fig.savefig(p, dpi=130)
    plt.close(fig)
    files.append(p)

    # voltage envelope
    fig, ax = plt.subplots(figsize=(12, 4.5))
    for j, e in enumerate(r["trainV"]["envelope"]):
        c = pal[j % len(pal)]
        ax.fill_between(e["x"], e["min"], e["max"], color=c, alpha=0.15)
        ax.plot(e["x"], e["min"], color=c, lw=1, label=f"{e['track']} min")
        ax.plot(e["x"], e["max"], color=c, lw=0.6, ls=":")
    ax.axhline(cfg["limits"]["V_min"], color="#b3261e", ls="--", lw=1, label="V min limit")
    ax.axhline(cfg["limits"]["V_max"], color="#b3261e", ls="-.", lw=1, label="V max limit")
    for ts in r["tss"]:
        ax.axvline(ts["ch"], color="#c05621" if ts["active"] else "#a0aec0", lw=0.6, ls="-" if ts["active"] else "--")
    ax.set_xlabel("Chainage (m)")
    ax.set_ylabel("Train voltage (V)")
    ax.set_title("Train voltage envelope (min / max by position)")
    ax.legend(ncol=6, fontsize=7)
    fig.tight_layout()
    p = os.path.join(outdir, "04_train_voltage.png")
    fig.savefig(p, dpi=130)
    plt.close(fig)
    files.append(p)

    # fault
    fig, ax = plt.subplots(figsize=(12, 4))
    for j, t in enumerate(flt["tracks"]):
        ax.plot(t["x"], np.array(t["If"]) / 1000, color=pal[j % len(pal)], label=f"{t['track']} total")
        ax.plot(t["x"], np.array(t["Ifeeder"]) / 1000, color=pal[j % len(pal)], ls=":", label=f"{t['track']} largest feeder")
    ax.axhline(cfg["limits"]["breaker_trip_A"] / 1000, color="#b3261e", ls="--", lw=1, label="Breaker trip setting")
    ax.set_xlabel("Fault chainage (m)")
    ax.set_ylabel("Fault current (kA)")
    ax.set_title("Steady-state DC short-circuit current along the line")
    ax.legend(fontsize=7)
    fig.tight_layout()
    p = os.path.join(outdir, "05_short_circuit.png")
    fig.savefig(p, dpi=130)
    plt.close(fig)
    files.append(p)
    return files


def to_jsonable(o):
    if isinstance(o, dict):
        return {k: to_jsonable(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [to_jsonable(v) for v in o]
    if isinstance(o, np.ndarray):
        return [to_jsonable(float(v)) for v in o.tolist()]
    if isinstance(o, (np.floating, float)):
        return float(o) if math.isfinite(o) else None
    if isinstance(o, np.integer):
        return int(o)
    if isinstance(o, np.bool_):
        return bool(o)
    return o


# ----------------------------------------------------------------------
# templates
# ----------------------------------------------------------------------
def make_templates(outdir):
    os.makedirs(outdir, exist_ok=True)
    cfg = E.defaults()
    A = cfg["alignment"]
    import xlsxwriter
    p1 = os.path.join(outdir, "Alignment_Template.xlsx")
    wb = xlsxwriter.Workbook(p1)
    hd = wb.add_format({"bold": True, "bg_color": "#1F3A4D", "font_color": "white", "border": 1})
    note = wb.add_format({"italic": True, "font_color": "#555555"})
    ws = wb.add_worksheet("Read me")
    ws.set_column(0, 0, 110)
    for i, t in enumerate([
        "Alignment input for the DC traction power model. Replace the sample rows with project data; keep the header rows.",
        "Chainages in metres (12345.6) or as km+m text (12+345.6). A header containing (km) means values are in km.",
        "Gradients: + = rising with increasing chainage (UP direction). Header '(‰)' = per mille, '(%)' = per cent, or text like '1 in 250 rise'.",
        "Sections may be given as From/To, as Start + Length, or as a point list (Chainage = start of section, runs to the next row).",
        "Speed limits: separate Up and Down columns, or one 'Speed (km/h)' column for both directions.",
        "Levels sheet (optional): Chainage + Rail level. Used for the long section, and to derive gradients if no gradient sheet is filled.",
        "Extra sheets and columns are ignored, so you can keep notes or drawing references alongside.",
    ]):
        ws.write(i, 0, t)
    sheets = [
        ("Stations", ["Station", "Chainage (m)", "Dwell (s)"], [[s["name"], s["ch"], s["dwell"]] for s in A["stations"]]),
        ("Gradients", ["From (m)", "To (m)", "Gradient (‰)"], [[g["from"], g["to"], g["g"]] for g in A["gradients"]]),
        ("Curves", ["From (m)", "To (m)", "Radius (m)"], [[c["from"], c["to"], c["R"]] for c in A["curves"]]),
        ("Speed limits", ["From (m)", "To (m)", "Up speed (km/h)", "Down speed (km/h)"], [[s["from"], s["to"], s["up"], s["dn"]] for s in A["speed"]]),
        ("Levels", ["Chainage (m)", "Rail level (m)"], []),
    ]
    for name, heads, rows in sheets:
        ws = wb.add_worksheet(name)
        ws.write_row(0, 0, heads, hd)
        for i, r in enumerate(rows):
            ws.write_row(i + 1, 0, r)
        ws.set_column(0, len(heads) - 1, 18)
        if not rows:
            ws.write(1, 0, "(optional - leave empty to use the gradient sheet)", note)
    wb.close()

    # Word train data sheet
    from docx import Document
    from docx.shared import Pt, Cm
    T = cfg["trains"][0]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Rolling stock data for traction power simulation", 1)
    doc.add_paragraph("Sample 6-car metro. Replace the values with the manufacturer's data. The model reads "
                      "'Parameter : value unit' lines, two-column parameter tables, curve tables with a Speed column, "
                      "Word charts, and graph images (axes calibrated in the digitiser).")
    doc.add_heading("1. General data", 2)
    rows = [("Number of cars", "6", ""), ("Number of traction motors", "16", ""), ("Train length", "132", "m"),
            ("Simulation mass (AW3, crush load)", str(T["mass_t"]), "t"), ("Rotating mass allowance", "8", "%"),
            ("Maximum operating speed", str(T["vmax"]), "km/h"), ("Starting acceleration", str(T["acc"]), "m/s²"),
            ("Service braking rate", str(T["dec"]), "m/s²"), ("Emergency braking rate", "1.3", "m/s²"),
            ("Auxiliary power per car", str(T["aux_kW_car"]), "kW per car"), ("Nominal line voltage", "750", "V DC"),
            ("Maximum regenerative braking power", str(T["regen_max_kW"]), "kW")]
    tb = doc.add_table(rows=1, cols=3)
    tb.style = "Light Grid Accent 1"
    for j, h in enumerate(["Parameter", "Value", "Unit"]):
        tb.rows[0].cells[j].text = h
    for r in rows:
        c = tb.add_row().cells
        for j, v in enumerate(r):
            c[j].text = v
    doc.add_paragraph()
    doc.add_paragraph(f"Running resistance (Davis): R = A + B·v + C·v² with A = {T['davis']['A']}, B = {T['davis']['B']}, "
                      f"C = {T['davis']['C']} (kN, v in km/h)")
    doc.add_heading("2. Tractive and electric braking effort (whole train, at wheel)", 2)
    be = dict((v, f) for v, f in T["be_curve"])
    eff = T["eff_curve"]
    tb = doc.add_table(rows=1, cols=4)
    tb.style = "Light Grid Accent 1"
    for j, h in enumerate(["Speed (km/h)", "Tractive effort (kN)", "Electric braking effort (kN)", "Drive efficiency (%)"]):
        tb.rows[0].cells[j].text = h
    import tps_engine as _e
    for v, f in T["te_curve"]:
        c = tb.add_row().cells
        c[0].text = f"{v:g}"
        c[1].text = f"{f:g}"
        c[2].text = f"{_e.interp(_e.clean_table(T['be_curve']), v):.1f}"
        c[3].text = f"{_e.interp(eff, v) * 100:.1f}"
    doc.add_heading("3. Tractive effort graph", 2)
    img = os.path.join(outdir, "_te_graph.png")
    _te_graph(T, img)
    doc.add_picture(img, width=Cm(15))
    doc.add_paragraph("Graph image: X axis 0 - 100 km/h, Y axis 0 - 350 kN (plot frame edges). Use these values when calibrating the digitiser.")
    p2 = os.path.join(outdir, "Train_Data_Template.docx")
    doc.save(p2)

    # second sample: graph only, to show image reading
    doc = Document()
    doc.add_heading("Train data - graph only sample", 1)
    doc.add_paragraph("Simulation mass (AW3): 260 t")
    doc.add_paragraph("Maximum operating speed: 80 km/h")
    doc.add_paragraph("Starting acceleration: 1.0 m/s2")
    doc.add_paragraph("Service braking rate: 1.0 m/s2")
    doc.add_paragraph("Auxiliary power: 240 kW per train")
    doc.add_paragraph("Number of cars: 6")
    doc.add_paragraph("Overall traction efficiency: 87 %")
    doc.add_paragraph("Tractive effort curve (X: 0-100 km/h, Y: 0-350 kN):")
    doc.add_picture(img, width=Cm(15))
    p3 = os.path.join(outdir, "Train_Data_Graph_Only.docx")
    doc.save(p3)
    os.remove(img)
    return [p1, p2, p3]


def _te_graph(T, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    v = np.linspace(0, 90, 181)
    tb = E.clean_table(T["te_curve"])
    f = [E.interp(tb, x) for x in v]
    fig, ax = plt.subplots(figsize=(7, 4))
    ax.plot(v, f, color="#d62728", lw=2.2, label="Tractive effort AW3")
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 350)
    ax.set_xlabel("Speed (km/h)")
    ax.set_ylabel("Tractive effort (kN)")
    ax.grid(alpha=0.35, color="#bbbbbb")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend()
    fig.tight_layout()
    fig.savefig(path, dpi=110)
    plt.close(fig)


# ----------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description="DC traction power simulation model", formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    ap.add_argument("--inputs", help="inputs JSON (exported from the HTML model or saved with --save-inputs)")
    ap.add_argument("--alignment", help="alignment Excel file")
    ap.add_argument("--train", help="train data Word file (.docx)")
    ap.add_argument("--train-index", type=int, default=0, help="which train type the Word file fills (0 = first)")
    ap.add_argument("--graph-axes", help="graph image axes: x_left,x_right,y_bottom,y_top (values at the plot frame edges)")
    ap.add_argument("--graph-use", default="te", choices=["te", "be", "eff", "lc"], help="what a graph image shows")
    ap.add_argument("--n-up", type=int, help="number of UP tracks")
    ap.add_argument("--n-down", type=int, help="number of DOWN tracks")
    ap.add_argument("--headway", type=float, help="set the same headway (s) on all tracks")
    ap.add_argument("--tss-count", type=int, help="number of TSS (regenerates the TSS list)")
    ap.add_argument("--tss-first", type=float, help="chainage of the first TSS (m)")
    ap.add_argument("--tss-spacing", type=float, help="distance between TSS (m)")
    ap.add_argument("--add-tss", action="append", help="add a TSS: name,chainage[,rect_kW[,xfmr_kVA]]  (repeatable)")
    ap.add_argument("--outage", action="append", default=[], help="TSS name to take out of service (repeatable)")
    ap.add_argument("--period", type=float, help="simulated period (s)")
    ap.add_argument("--design-check", action="store_true", help="also run every single-TSS outage")
    ap.add_argument("--out", default="results", help="output folder")
    ap.add_argument("--save-inputs", help="write the final inputs to this JSON file")
    ap.add_argument("--make-templates", action="store_true", help="write input templates to the output folder and stop")
    ap.add_argument("--no-timeseries", action="store_true", help="skip the per-second sheet (faster Excel)")
    ap.add_argument("--no-reports", action="store_true", help="skip the four Word reports (DC, short circuit, harmonics, EMC)")
    a = ap.parse_args(argv)
    os.makedirs(a.out, exist_ok=True)
    if a.make_templates:
        for p in make_templates(a.out):
            print("wrote", p)
        return 0
    log = []
    cfg = load_inputs(a.inputs)
    cfg = apply_cli(cfg, a, log)
    for line in log:
        print(line)
    if a.save_inputs:
        with open(a.save_inputs, "w", encoding="utf-8") as f:
            json.dump(cfg, f, indent=1)
        print("inputs saved to", a.save_inputs)
    t0 = time.time()
    print(f"Simulating {cfg['sim']['period']:.0f} s, {len(cfg['tracks'])} tracks, {len(cfg['tss'])} TSS ...")
    r = E.simulate(cfg, outage=a.outage)
    flt = E.fault_scan(cfg, outage=a.outage)
    st = E.studies(cfg, r, flt)
    r["checks"] = st["checks"]
    print(f"  done in {time.time() - t0:.1f} s (simulation, short circuit, harmonics, EMC)")
    dc = None
    if a.design_check:
        dc = E.design_check(cfg, on_progress=lambda i, n, nm: print(f"  design check {i}/{n}: {nm}"))
    xl = os.path.join(a.out, "TPS_Results.xlsx")
    write_excel(xl, cfg, r, flt, dc, st, timeseries=not a.no_timeseries)
    pngs = write_charts(a.out, cfg, r, flt)
    summ = {"scenario": r["scenario"], "trainV": r["trainV"], "energy": r["energy"], "tss": r["tss"], "tracks": r["tracks"],
            "checks": r["checks"], "fault": [{k: t[k] for k in ("track", "max", "min", "minFeeder", "minDidtFeeder", "maxTtrip", "maxUrail")} for t in flt["tracks"]],
            "fault_tss": flt["tssAC"],
            "harmonics": [{k: t.get(k) for k in ("name", "THDv95", "THDvmax", "THDi_rated", "mu_rated_deg")} for t in st["harm"]["tss"] if t["active"]],
            "emc": {k: st["emc"][k] for k in ("max1", "max300", "maxMean", "stray_mA_m", "faultUrail", "field", "Mprime_uH_m")},
            "design_check": [{k: v for k, v in x.items() if k != "checks"} for x in dc] if dc else None}
    with open(os.path.join(a.out, "TPS_Results_summary.json"), "w", encoding="utf-8") as f:
        json.dump(to_jsonable(summ), f, indent=1)
    ch = r["checks"]
    fails = [c for c in ch if not c["pass"]]
    print(f"\nScenario: {r['scenario']}   checks: {len(ch) - len(fails)} pass, {len(fails)} fail")
    print(f"Train voltage min {r['trainV']['min']:.0f} V / max {r['trainV']['max']:.0f} V, energy from TSS {r['energy']['tss_kWh']:.0f} kWh")
    print(f"{'TSS':<10}{'1-h RMS A':>11}{'% rating':>10}{'15-min A':>10}{'1-min A':>9}{'Peak A':>9}{'Xfmr kVA':>10}")
    for t in r["tss"]:
        if t["active"]:
            print(f"{t['name']:<10}{t['I1h']:>11.0f}{t['I1h'] / t['Ir'] * 100:>9.0f}%{t['I15']:>10.0f}{t['I1']:>9.0f}{t['Ipk']:>9.0f}{t['S1h']:>10.0f}")
        else:
            print(f"{t['name']:<10}{'out of service':>20}")
    for c in fails[:12]:
        print(f"  FAIL  {c['group']}: {c['item']} at {c['where']}: {c['value']:.0f} {c['unit']} vs {c['limit']:.0f} ({c['ratio'] * 100:.0f} %)")
    if dc:
        print("\nDesign check:")
        for x in dc:
            print(f"  {x['scenario']:<22} Vmin {x['vmin']:.0f} V  rect {x['rect'] * 100:.0f} %  xfmr {x['xfmr'] * 100:.0f} %  failures {x['nFail']}")
    H, em = st["harm"], st["emc"]
    wt = max([t for t in H["tss"] if t["active"]] + ([H["common"]] if H["common"] else []), key=lambda t: t["THDv95"])
    print(f"Short circuit: max {max(t['max'] for t in flt['tracks']) / 1000:.1f} kA, min feeder {min(t['minFeeder'] for t in flt['tracks']) / 1000:.1f} kA, "
          f"min feeder di/dt {min(t['minDidtFeeder'] for t in flt['tracks']):.0f} A/ms")
    print(f"Harmonics: worst voltage THD ({H['percentile']:g} %) {wt['THDv95']:.2f} % at {wt['name']} (limit {H['thd_limit']} %)")
    print(f"EMC: rail potential {em['max1']:.0f} V peak / {em['max300']:.0f} V 300 s, stray {em['stray_mA_m']:.2f} mA/m, "
          f"field {em['field'][0]['Bpk_uT']:.0f} uT at {em['field'][0]['d']:g} m" if em["field"] else "")
    if not a.no_reports:
        for p in RP.write_all(cfg, r, flt, st, a.out, dc):
            print("wrote", p)
    print("\nwrote", xl)
    for p in pngs:
        print("wrote", p)
    return 0


if __name__ == "__main__":
    sys.exit(main())
