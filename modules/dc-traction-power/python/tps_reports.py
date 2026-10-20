"""
Word reports for the DC traction power model
=============================================
Four reports with the same structure as the reports of the HTML model:

  DC  - DC traction power simulation        SC  - short circuit study
  HM  - harmonic study                      EMC - EMC / EMI study

Each report is first built as a list of blocks (headings, text, key/value
tables, tables, charts, bullets) and then written with python-docx.
"""
from __future__ import annotations

import datetime
import io
import math
import os

import numpy as np

import tps_engine as E

PC = {"ink": "#1c2631", "ink2": "#4d5a67", "grid": "#e1e4dd", "up": "#1d6c88", "dn": "#7a4aa0", "acc": "#b1531b",
      "fail": "#b3261e", "pass": "#1b7442", "mute": "#8a949e"}
PPAL = [PC["up"], PC["dn"], PC["acc"], PC["pass"], "#9a6400", "#5b6f8a", "#b04a7a", "#3d8f8a"]
REPORTS = [("dc", "DC", "DC Traction Power Simulation Report"), ("sc", "SC", "Short Circuit Study Report"),
           ("hm", "HM", "Harmonic Study Report"), ("emc", "EMC", "EMC / EMI Study Report")]


def fmt(v, d=0):
    if v is None:
        return "–"
    try:
        v = float(v)
    except (TypeError, ValueError):
        return str(v)
    if not math.isfinite(v):
        return "–"
    return f"{v:,.{d}f}"


def f0(v):
    return fmt(v, 0)


def f1(v):
    return fmt(v, 1)


def f2(v):
    return fmt(v, 2)


def mmss(s):
    s = int(math.floor(s + 0.5))
    return f"{s // 60}:{s % 60:02d}"


# ----------------------------------------------------------------------
# charts
# ----------------------------------------------------------------------
def _fig(h=4.0):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 9, "axes.grid": True, "grid.color": PC["grid"], "axes.spines.top": False,
                         "axes.spines.right": False, "axes.edgecolor": PC["ink2"], "text.color": PC["ink"],
                         "axes.labelcolor": PC["ink2"], "xtick.color": PC["ink2"], "ytick.color": PC["ink2"]})
    fig, ax = plt.subplots(figsize=(10, h))
    return plt, fig, ax


def _png(plt, fig, caption):
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    return {"t": "chart", "png": buf.getvalue(), "caption": caption}


def _vlines(ax, cfg):
    for t in cfg["tss"]:
        ax.axvline(float(t["ch"]), color=PC["mute"], lw=0.7, ls=":")
        ax.text(float(t["ch"]), 1.0, t["name"], transform=ax.get_xaxis_transform(), fontsize=7, ha="center", va="bottom", color=PC["ink2"])


# ----------------------------------------------------------------------
# blocks
# ----------------------------------------------------------------------
def _check_table(r, groups):
    rows = [c for c in r["checks"] if c["group"] in groups]
    dec = lambda u: 2 if u in ("kA", "%", "mT", "mV", "mA/m", "S/km") else 0
    return {"t": "table", "head": ["Result", "Group", "Check", "Location", "Value", "Limit", "Unit", "Use %"],
            "rows": [["PASS" if c["pass"] else "FAIL", c["group"], c["item"], c["where"], fmt(c["value"], dec(c["unit"])),
                      ("≥ " if c["kind"] == "min" else "≤ ") + fmt(c["limit"], dec(c["unit"])), c["unit"], f0(c["ratio"] * 100)] for c in rows]}


def _verdict(r, groups):
    rows = [c for c in r["checks"] if c["group"] in groups]
    f = [c for c in rows if not c["pass"]]
    if f:
        c = f[0]
        return (f"{len(f)} of {len(rows)} checks fail. The first is: {c['group']}, {c['item']} at {c['where']} "
                f"({fmt(c['value'], 1)} {c['unit']} against {fmt(c['limit'], 1)}).")
    return f"All {len(rows)} checks pass for this scenario."


def _line_facts(cfg, r):
    A = cfg["alignment"]
    ext = r["extent"]
    gmax = max([abs(float(g["g"])) for g in A["gradients"]] + [0])
    rmin = min([float(c["R"]) for c in A["curves"] if float(c["R"])] + [1e9])
    return [["Line length (model extent)", f0(ext[1] - ext[0]) + " m"], ["Stations", str(len(A["stations"]))],
            ["Tracks", ", ".join(f"{t['name']} ({t['dir']})" for t in cfg["tracks"] if t.get("enabled", True) is not False)],
            ["Steepest gradient", f1(gmax) + " ‰"], ["Smallest curve radius", f0(rmin) + " m" if rmin < 1e9 else "none"],
            ["Traction substations", f"{len(cfg['tss'])} ({sum(1 for t in cfg['tss'] if t.get('in_service', True) is not False)} in service)"]]


def _sc_extremes(flt):
    near = far = None
    for t in flt["tracks"]:
        for i, x in enumerate(t["x"]):
            c = {"track": t["track"], "x": x, "If": t["If"][i], "Ifd": t["Ifeeder"][i], "tau": t["tau_ms"][i],
                 "didtF": t["didtFeeder"][i], "U": t["Urail"][i]}
            if near is None or c["If"] > near["If"]:
                near = c
            if far is None or c["Ifd"] < far["Ifd"]:
                far = c
    return near, far


def _harm_entities(harm):
    e = [t for t in harm["tss"] if t["active"]]
    return ([harm["common"]] + e) if harm["common"] else e


def build_report(rid, cfg, r, flt, st, dc=None):
    code, title = {k: (c, t) for k, c, t in REPORTS}[rid]
    ref = cfg["project"].get("ref") or "TPS"
    rp = {"title": title, "docref": f"{ref}-{code}", "file": f"{ref}_{code}_Report".replace(" ", "_")}
    A = cfg["alignment"]
    B = [{"t": "title", "title": title, "sub": cfg["project"]["name"], "meta": [
        ["Document", rp["docref"]], ["System", cfg["project"].get("system", "")], ["Scenario", r["scenario"]],
        ["Date", datetime.date.today().isoformat()], ["Alignment data", A.get("source", "")],
        ["Train data", ", ".join(t.get("source") or t.get("name", "") for t in cfg["trains"])], ["Calculation", r["engine"]]]}]
    El, L = cfg["electrical"], cfg["limits"]
    harm, em = st["harm"], st["emc"]
    plt = None
    if rid == "dc":
        groups = ["Rectifier", "Transformer", "Feeder", "Conductor rail", "Train voltage"]
        B += [{"t": "h1", "text": "Summary"},
              {"t": "p", "text": f"This report gives the DC traction power simulation of {cfg['project']['name']} for the scenario "
                                 f"\"{r['scenario']}\" over {f0(r['period'])} s of peak service. " + _verdict(r, groups)},
              {"t": "kv", "rows": [["Minimum train voltage", f"{f0(r['trainV']['min'])} V (limit {f0(L['V_min'])} V)"],
                                   ["Maximum train voltage", f"{f0(r['trainV']['max'])} V (limit {f0(L['V_max'])} V)"],
                                   ["Mean useful voltage", f"{f0(r['trainV']['mean_useful'])} V"],
                                   ["Energy from substations", f"{f0(r['energy']['tss_kWh'])} kWh"],
                                   ["Regenerated energy returned", f"{f0(r['energy']['regen_out_kWh'])} kWh"],
                                   ["Network losses", f"{f0(r['energy']['losses_kWh'])} kWh ({f1(r['energy']['losses_kWh'] / max(r['energy']['tss_kWh'], 1) * 100)} %)"]]},
              {"t": "h1", "text": "Scope and method"},
              {"t": "p", "text": f"Each train service is run over the alignment on a {f0(E.num(cfg['sim'].get('dx'), 5))} m distance grid with tractive and "
                                 "braking effort, Davis resistance, gradient and curve resistance, speed restrictions held for the train length and the "
                                 "station dwell times. The trains are placed on the tracks at their headways and the DC network (substations, feeders, "
                                 f"conductor and return rails, trains) is solved every {f0(r['dt'])} s with diode substations, constant-power trains, "
                                 "low-voltage cut-back and the regeneration voltage limit. Ratings are checked against rolling 1-minute, 15-minute and 1-hour RMS values."},
              {"t": "h1", "text": "Input data"}, {"t": "h2", "text": "Line and alignment"}, {"t": "kv", "rows": _line_facts(cfg, r)},
              {"t": "h2", "text": "Tracks and service"},
              {"t": "table", "head": ["Track", "Direction", "Headway s", "Trains/h", "Train type", "Skip stations", "Trip min"],
               "rows": [[t["name"], t["dir"], f0(t["headway"]), f1(t["tph"]),
                         cfg["trains"][int(E.num(next((c for c in cfg["tracks"] if c["name"] == t["name"]), {}).get("train"), 0))]["name"],
                         next((c for c in cfg["tracks"] if c["name"] == t["name"]), {}).get("skip") or "–", f1(t["trip"] / 60)] for t in r["tracks"]]},
              {"t": "h2", "text": "Traction substations"},
              {"t": "table", "head": ["TSS", "Chainage m", "Rectifier kW", "Transformer kVA", "No-load V", "Source R mΩ", "Inverter kW", "Status"],
               "rows": [[t["name"], f0(t["ch"]), f0(t["rect_kW"]), f0(t["xfmr_kVA"]), f0(t["V0"]), f2(t["Req"] * 1000), f0(t["inv_kW"]),
                         "in service" if t["active"] else "out"] for t in r["tss"]]},
              {"t": "h2", "text": "Rolling stock"},
              {"t": "table", "head": ["Train", "Cars", "Mass t", "Length m", "Vmax km/h", "Acc m/s²", "Brake m/s²", "Aux kW", "Davis A/B/C"],
               "rows": [[t["name"], t.get("cars"), t.get("mass_t"), t.get("length_m"), t.get("vmax"), t.get("acc"), t.get("dec"),
                         f0(E.num(t.get("aux_kW_car"), 0) * E.num(t.get("cars"), 1)),
                         f"{t['davis']['A']} / {t['davis']['B']} / {t['davis']['C']}"] for t in cfg["trains"]]}]
        T0 = cfg["trains"][0]
        mm = max(1, E.num(T0.get("n_motors"), 1)) if T0.get("te_basis") == "motor" else 1
        plt, fig, ax = _fig(3.6)
        te, be = E.clean_table(T0["te_curve"]), E.clean_table(T0["be_curve"])
        ax.plot([p[0] for p in te], [p[1] * mm for p in te], color=PC["up"], lw=2, label="Tractive effort")
        ax.plot([p[0] for p in be], [p[1] * mm for p in be], color=PC["dn"], lw=2, ls="--", label="Electric braking effort")
        ax.set_xlabel("Speed (km/h)"); ax.set_ylabel("Effort per train (kN)"); ax.legend()
        B.append(_png(plt, fig, f"Effort curves of {T0['name']}"))
        B += [{"t": "h2", "text": "Electrical parameters and limits"},
              {"t": "kv", "rows": [["Nominal / no-load voltage", f"{El['V_nom']} V / {El['V0']} V"],
                                   ["Conductor + return resistance", f"{El['r_pos']} + {El['r_ret']} Ω/km per track"],
                                   ["Feeder resistance", f"{El['R_feeder']} Ω"], ["Regeneration limit", f"{El['V_regen_start']} → {El['V_regen_max']} V"],
                                   ["Traction cut-back", f"{El['V_cut_start']} → {El['V_cut_zero']} V"],
                                   ["Train voltage limits", f"{L['V_min']} V min, {L['V_max']} V max, {L['U_mean_useful']} V mean useful"],
                                   ["Conductor rail / feeder ratings", f"{L['conductor_A']} A / {L['feeder_A']} A"]]},
              {"t": "h1", "text": "Results"}, {"t": "h2", "text": "Train runs"}]
        for k, run in r["runs"].items():
            tr = ", ".join(t["name"] for t in r["tracks"] if t["key"] == k)
            B.append({"t": "p", "text": f"{tr} · {run['train']} · {run['dir']}: trip {mmss(run['T'])} min:s, top speed {f1(run['vmax_kmh'])} km/h, "
                                        f"traction {f1(run['E_trac_kWh'])} kWh, auxiliaries {f1(run['E_aux_kWh'])} kWh, regeneration available {f1(run['E_regen_kWh'])} kWh."})
        r0 = next(iter(r["runs"].values()))
        plt, fig, ax = _fig(3.4)
        ax.plot(r0["x"], r0["v"], color=PC["up"], lw=1.4)
        for s_ in r0["stations"]:
            ax.axvline(s_["ch"], color=PC["mute"], lw=0.5, ls=":")
        ax.set_xlabel("Chainage (m)"); ax.set_ylabel("Speed (km/h)")
        B.append(_png(plt, fig, f"Speed profile, {r0['train']} {r0['dir']}"))
        B.append({"t": "table", "head": ["Station", "Chainage m", "Arrive", "Depart"], "rows": [[s_["name"], f0(s_["ch"]), mmss(s_["arr"]), mmss(s_["dep"])] for s_ in r0["stations"]]})
        d = L["contingency"] if r["contingency"] else L["normal"]
        Vn = E.num(El.get("V_nom"), 750)
        B += [{"t": "h2", "text": "Substation loading and equipment sizing"},
              {"t": "table", "head": ["TSS", "Rated A", "Peak A", "1-min A", "15-min A", "1-h A", "1-h %", "Needed kW", "Xfmr 1-h kVA", "Needed kVA", "Energy kWh"],
               "rows": [[t["name"], f0(t["Ir"]), f0(t["Ipk"]), f0(t["I1"]), f0(t["I15"]), f0(t["I1h"]), f0(t["I1h"] / t["Ir"] * 100),
                         f0(max(t["I1h"] / d["h1"], t["I15"] / d["m15"], t["I1"] / d["m1"]) * Vn / 1000), f0(t["S1h"]),
                         f0(max(t["S1h"] / d["h1"], t["S15"] / d["m15"], t["S1"] / d["m1"])), f0(t["E_kWh"])] if t["active"]
                        else [t["name"], f0(t["Ir"]), "out"] + [""] * 8 for t in r["tss"]]}]
        top = sorted([(t["I1h"] if t["active"] else -1, i) for i, t in enumerate(r["tss"])], reverse=True)[:3]
        plt, fig, ax = _fig(3.6)
        tm = np.arange(r["ts"]["tssI"].shape[1]) * r["dt"] / 60
        for q, (_, i) in enumerate(top):
            ax.plot(tm, r["ts"]["tssI"][i], lw=0.7, color=PPAL[q], label=r["tss"][i]["name"])
        ax.set_xlabel("Time (min)"); ax.set_ylabel("Rectifier current (A)"); ax.legend()
        B.append(_png(plt, fig, "Currents of the three most loaded substations"))
        B.append({"t": "h2", "text": "Train voltage"})
        plt, fig, ax = _fig(3.8)
        for q, e in enumerate(r["trainV"]["envelope"]):
            ax.plot(e["x"], e["min"], color=PPAL[q], lw=1.2, label=f"{e['track']} lowest")
            ax.plot(e["x"], e["max"], color=PPAL[q], lw=0.8, ls=":", label=f"{e['track']} highest")
        ax.axhline(L["V_min"], color=PC["fail"], ls="--", lw=1, label="Minimum limit")
        _vlines(ax, cfg)
        ax.set_xlabel("Chainage (m)"); ax.set_ylabel("Train voltage (V)"); ax.legend(ncol=5, fontsize=7)
        B.append(_png(plt, fig, "Lowest and highest train voltage by position"))
        B += [{"t": "h2", "text": "Feeders and conductor rail"},
              {"t": "table", "head": ["Feeder", "Track", "Peak A", "1-h RMS A", "% of cable rating"],
               "rows": [[f["tss"], f["track"], f0(f["Ipk"]), f0(f["I1h"]), f0(f["I1h"] / L["feeder_A"] * 100)]
                        for f in sorted(r["feeders"], key=lambda f: -f["I1h"])[:12]]}]
        if dc:
            B += [{"t": "h1", "text": "Design check, all cases"},
                  {"t": "p", "text": f"{sum(1 for x in dc if x['nFail'])} of {len(dc)} cases have failing checks. Outage cases use the contingency duty multipliers."},
                  {"t": "table", "head": ["Case", "Min V", "Max V", "Rect %", "Xfmr %", "Rail %", "Feeder %", "THD %", "Fails", "First failure"],
                   "rows": [[x["scenario"], f0(x["vmin"]), f0(x["vmax"]), f0(x["rect"] * 100), f0(x["xfmr"] * 100), f0(x["cond"] * 100),
                             f0(x["feeder"] * 100), f2(x["thd"]), str(x["nFail"]), x["firstFail"]] for x in dc]}]
        B += [{"t": "h1", "text": "Compliance checks"}, _check_table(r, groups)]
    elif rid == "sc":
        S, AC = cfg["sc"], cfg["ac"]
        groups = ["Short circuit"]
        near, far = _sc_extremes(flt)
        fr = [c for c in r["checks"] if c["group"] == "EMC" and "fault" in c["item"]]
        B += [{"t": "h1", "text": "Summary"},
              {"t": "p", "text": f"Short circuit study for {cfg['project']['name']}, scenario \"{r['scenario']}\". " + _verdict(r, groups) +
                                 (f" Rail potential during faults: {f0(fr[0]['value'])} V against {f0(fr[0]['limit'])} V ({'PASS' if fr[0]['pass'] else 'FAIL'})." if fr else "")},
              {"t": "kv", "rows": [["Maximum DC fault current", f"{f1(max(t['max'] for t in flt['tracks']) / 1000)} kA at {near['track']} {f0(near['x'])} m"],
                                   ["Minimum feeder fault current", f"{f1(min(t['minFeeder'] for t in flt['tracks']) / 1000)} kA at {far['track']} {f0(far['x'])} m"],
                                   ["Breaker trip setting / breaking capacity", f"{f1(L['breaker_trip_A'] / 1000)} kA / {f0(L['breaking_kA'])} kA"],
                                   ["Minimum feeder di/dt", f"{f0(min(t['minDidtFeeder'] for t in flt['tracks']))} A/ms (setting {S['didt_setting_A_ms']} A/ms)"]]},
              {"t": "h1", "text": "Method"},
              {"t": "p", "text": "AC side: three-phase fault currents to IEC 60909 at the HV busbar of each substation (network fault level, X/R and "
                                 "voltage factor c) and at the rectifier valve windings (network plus transformer impedance per bridge), with peak "
                                 "factor κ = 1.02 + 0.98·e^(−3R/X)."},
              {"t": "p", "text": f"DC side: a bolted fault (fault resistance {cfg['sim']['R_fault']} Ω) is placed every {cfg['sim']['fault_step']} m along every "
                                 "track. The resistance network gives the sustained current; the inductance network (rectifier commutation inductance "
                                 f"2·Lc per bridge plus {AC['L_smooth_mH']} mH smoothing reactor, feeder {S['L_feeder_mH']} mH, loop {S['L_line_mH_km']} mH/km) gives "
                                 "the initial rate of rise. The current rises as i(t) = Iss·(1 − e^(−t/τ)), τ = Iss ÷ (di/dt). Breaker detection uses the "
                                 "largest single-feeder contribution. Rail potential at the fault = If·½√(r/g)·(1 − e^(−γd))."},
              {"t": "h1", "text": "Input data"},
              {"t": "kv", "rows": [["Supply", f"{AC['U_hv_kV']} kV, fault level {AC['S_sc_MVA']} MVA, X/R {AC['XR']}, c = {AC['c_max']}"],
                                   ["Rectifier transformer", f"uk {AC['uk_pct']} %, load losses {AC['pcu_pct']} %, {AC['pulse']}-pulse"],
                                   ["No-load DC voltage", f"{El['V0']} V"], ["Loop / feeder inductance", f"{S['L_line_mH_km']} mH/km / {S['L_feeder_mH']} mH"],
                                   ["Breaker trip / breaking / di/dt setting", f"{L['breaker_trip_A']} A / {L['breaking_kA']} kA / {S['didt_setting_A_ms']} A/ms"],
                                   ["Fault clearing time", f"{S['t_clear_s']} s"]]},
              {"t": "h1", "text": "AC fault levels and rectifier data"},
              {"t": "table", "head": ["TSS", "HV Ik″ kA", "HV ip kA", "Valve Ik″ kA", "Valve ip kA", "Xc mΩ", "R used mΩ", "R from AC mΩ", "L mH", "DC term. kA", "τ ms"],
               "rows": [[a["name"], f2(a["IkHV_kA"]), f2(a["ipHV_kA"]), f1(a["IkLV_kA"]), f1(a["ipLV_kA"]), f2(a["Xc_mohm"]), f2(a["Req_mohm"]),
                         f2(a["ReqAC_mohm"]), fmt(a["Ldc_mH"], 3), f1(a["IkD_kA"]), f2(a["tauD_ms"])] for a in flt["tssAC"]]},
              {"t": "h1", "text": "DC fault currents along the line"},
              {"t": "table", "head": ["Track", "Max kA", "Min kA", "Min feeder kA", "Min feeder di/dt A/ms", "Longest time to trip ms", "Max rail potential V"],
               "rows": [[t["track"], f1(t["max"] / 1000), f1(t["min"] / 1000), f1(t["minFeeder"] / 1000), f0(t["minDidtFeeder"]),
                         (f1(t["maxTtrip"]) + (f" ({t['notTripped']} not reached)" if t["notTripped"] else "")) if t["maxTtrip"] is not None else "not reached",
                         f0(t["maxUrail"])] for t in flt["tracks"]]}]
        plt, fig, ax = _fig(3.8)
        for q, t in enumerate(flt["tracks"]):
            ax.plot(t["x"], np.array(t["If"]) / 1000, color=PPAL[q], lw=1.5, label=f"{t['track']} total")
            ax.plot(t["x"], np.array(t["Ifeeder"]) / 1000, color=PPAL[q], lw=1, ls="--", label=f"{t['track']} largest feeder")
        ax.axhline(L["breaker_trip_A"] / 1000, color=PC["fail"], ls="--", lw=1, label="Trip setting")
        _vlines(ax, cfg); ax.set_ylim(bottom=0)
        ax.set_xlabel("Fault chainage (m)"); ax.set_ylabel("Sustained fault current (kA)"); ax.legend(fontsize=7, ncol=3)
        B.append(_png(plt, fig, "Sustained fault current along each track"))
        plt, fig, ax = _fig(3.4)
        for q, t in enumerate(flt["tracks"]):
            ax.semilogy(t["x"], t["didtFeeder"], color=PPAL[q], lw=1.4, label=f"{t['track']} largest feeder")
        ax.axhline(E.num(S.get("didt_setting_A_ms"), 40), color=PC["fail"], ls="--", lw=1, label="di/dt setting")
        _vlines(ax, cfg); ax.set_xlabel("Fault chainage (m)"); ax.set_ylabel("Initial di/dt (A/ms)"); ax.legend(fontsize=7)
        B.append(_png(plt, fig, "Rate of rise seen by the largest feeder"))
        tmax = max(5 * max(near["tau"], far["tau"]), 20)
        tt = np.linspace(0, tmax, 200)
        tauf = far["Ifd"] / max(far["didtF"], 1e-6)
        plt, fig, ax = _fig(3.2)
        ax.plot(tt, near["If"] * (1 - np.exp(-tt / max(near["tau"], 1e-9))) / 1000, color=PC["acc"], lw=1.6, label=f"Nearest fault ({f0(near['x'])} m)")
        ax.plot(tt, far["Ifd"] * (1 - np.exp(-tt / max(tauf, 1e-9))) / 1000, color=PC["up"], lw=1.6, label=f"Remote fault, feeder ({f0(far['x'])} m)")
        ax.axhline(L["breaker_trip_A"] / 1000, color=PC["fail"], ls="--", lw=1, label="Trip setting")
        ax.set_xlabel("Time after fault (ms)"); ax.set_ylabel("Current (kA)"); ax.legend(fontsize=7); ax.set_ylim(bottom=0)
        B.append(_png(plt, fig, "Fault current rise"))
        B += [{"t": "h1", "text": "Rail potential during faults"},
              {"t": "p", "text": f"Highest rail potential at a fault {f0(max(t['maxUrail'] for t in flt['tracks']))} V. Touch-voltage limit for {S['t_clear_s']} s: "
                                 f"{f0(E.touch_limit(cfg, E.num(S.get('t_clear_s'), 0.06)))} V. Rail–earth conductance {cfg['emc']['g_track_S_km']} S/km per track; "
                                 "the touch voltage is conservatively taken equal to the rail potential."},
              {"t": "h1", "text": "Compliance checks"}, _check_table(r, groups)]
        if fr:
            B.append({"t": "table", "head": ["Result", "Check", "Value V", "Limit V"], "rows": [["PASS" if c["pass"] else "FAIL", c["item"], f0(c["value"]), f0(c["limit"])] for c in fr]})
    elif rid == "hm":
        H, AC = harm, cfg["ac"]
        ents = _harm_entities(H)
        w = max(ents, key=lambda t: t["THDv95"])
        groups = ["Harmonics"]
        B += [{"t": "h1", "text": "Summary"},
              {"t": "p", "text": f"Harmonic study of the {H['pulse']}-pulse traction rectifiers for {cfg['project']['name']}, scenario \"{r['scenario']}\". "
                                 f"The highest voltage THD ({fmt(H['percentile'])} % value) is {f2(w['THDv95'])} % at {w['name']} against a limit of {f1(H['thd_limit'])} %. "
                                 + _verdict(r, groups)},
              {"t": "h1", "text": "Method"},
              {"t": "p", "text": "For every second of the simulated substation current the commutation overlap of each diode bridge is found from "
                                 "1 − cos μ = √2·Xc·Id/Uv. Harmonic currents follow the overlap formula for uncontrolled bridges, "
                                 "Ih/I1 = √(A² + B² − 2AB·cos μ) ÷ (h·(1 − cos μ)), A = sin((h−1)μ/2)/(h−1), B = sin((h+1)μ/2)/(h+1). "
                                 f"Characteristic orders kp ± 1 are kept in full; other 6k ± 1 orders at {cfg['harm']['residual_pct']} %. "
                                 "Harmonic voltages at the HV supply are Ih·h·Xnet. " +
                                 ("All substations share one supply bus and are added with the IEC 61000-3-6 summation exponents. " if cfg["harm"].get("pcc") == "common"
                                  else "Each substation has its own supply point. ") +
                                 f"The {fmt(H['percentile'])} % value over the period is compared with the limits."},
              {"t": "h1", "text": "Input data"},
              {"t": "kv", "rows": [["Supply", f"{AC['U_hv_kV']} kV, fault level {AC['S_sc_MVA']} MVA, X/R {AC['XR']}"],
                                   ["Rectifier", f"{AC['pulse']}-pulse, transformer uk {AC['uk_pct']} %"], ["Highest order", str(H["orders"][-1])],
                                   ["Limits", f"{H['std_name']}, THD {H['thd_limit']} %"]]},
              {"t": "table", "head": ["Order h"] + [str(h) for h, l in zip(H["orders"], H["limits"]) if l is not None],
               "rows": [["Limit %"] + [f2(l) for l in H["limits"] if l is not None]]},
              {"t": "h1", "text": "Results"},
              {"t": "table", "head": ["Supply point", "Overlap °", "HV I1 A", "I THD rated %", f"V THD {fmt(H['percentile'])}% %", "V THD max %", "Result"],
               "rows": [[t["name"], f1(t.get("mu_rated_deg")), f1(t.get("I1_rated")), f1(t.get("THDi_rated")), f2(t["THDv95"]), f2(t["THDvmax"]),
                         "PASS" if t["THDv95"] <= H["thd_limit"] and all(l is None or v <= l for v, l in zip(t["Vh95"], H["limits"])) else "FAIL"] for t in ents]}]
        plt, fig, ax = _fig(3.8)
        xs = np.arange(len(H["orders"]))
        ax.bar(xs - 0.2, w["Vh95"], 0.4, color=[PC["fail"] if l is not None and v > l else PC["up"] for v, l in zip(w["Vh95"], H["limits"])], label=f"{fmt(H['percentile'])} % value")
        ax.bar(xs + 0.2, w["Vhmax"], 0.4, color="#cfd6dc", label="Maximum")
        ax.step(xs, [l if l is not None else np.nan for l in H["limits"]], where="mid", color=PC["fail"], lw=1.2, label="Limit")
        ax.set_xticks(xs, [str(h) for h in H["orders"]]); ax.set_xlabel("Harmonic order"); ax.set_ylabel("% of nominal voltage"); ax.legend()
        B.append(_png(plt, fig, f"Harmonic voltage spectrum at {w['name']}"))
        B.append({"t": "table", "head": ["h", "Type", f"Vh {fmt(H['percentile'])}% %", "Vh max %", "Limit %", "Result"],
                  "rows": [[str(h), "characteristic" if H["characteristic"][j] else "residual", fmt(w["Vh95"][j], 3), fmt(w["Vhmax"][j], 3),
                            f2(H["limits"][j]) if H["limits"][j] is not None else "–",
                            ("PASS" if w["Vh95"][j] <= H["limits"][j] else "FAIL") if H["limits"][j] is not None else ""] for j, h in enumerate(H["orders"])]})
        plt, fig, ax = _fig(3.0)
        ax.plot(np.arange(len(w["ts"]["THDv"])) * r["dt"] / 60, w["ts"]["THDv"], color=PC["up"], lw=0.7)
        ax.axhline(H["thd_limit"], color=PC["fail"], ls="--", lw=1)
        ax.set_ylim(bottom=0); ax.set_xlabel("Time (min)"); ax.set_ylabel("Voltage THD (%)")
        B.append(_png(plt, fig, f"Voltage THD through the period at {w['name']}"))
        B += [{"t": "h2", "text": "DC side ripple"},
              {"t": "table", "head": ["TSS", "f Hz", "Ripple V", "Ripple A"],
               "rows": [[t["name"] if k == 0 else "", f0(d["f"]), f2(d["V"]), f2(d["I"])] for t in H["tss"] if t["active"] for k, d in enumerate(t["dc"])]},
              {"t": "h1", "text": "Compliance checks"}, _check_table(r, groups)]
        if any(c["group"] == "Harmonics" and not c["pass"] for c in r["checks"]):
            B.append({"t": "p", "text": "Mitigation options to consider: higher pulse number by phase-shifted transformers between substations, tuned filters "
                                        "at the supply point, or a supply point with a higher fault level. Agree limits and background distortion with the supply authority."})
    elif rid == "emc":
        M, Lm = cfg["emc"], em["limits"]
        groups = ["EMC"]
        B += [{"t": "h1", "text": "Summary"}, {"t": "p", "text": f"EMC / EMI study for {cfg['project']['name']}, scenario \"{r['scenario']}\". " + _verdict(r, groups)},
              {"t": "kv", "rows": [["Highest rail potential", f"{f0(em['max1'])} V at {f0(em['at1'])} m (limit {f0(Lm['s1'])} V)"],
                                   ["Highest 300 s average", f"{f0(em['max300'])} V (limit {f0(Lm['s300'])} V)"],
                                   ["Stray current", f"{f2(em['stray_mA_m'])} mA/m (limit {M['stray_limit_mA_m']} mA/m)"],
                                   [f"Static field at {em['field'][0]['d'] if em['field'] else '–'} m", f"{f0(em['field'][0]['Bpk_uT'])} µT at peak current" if em["field"] else "–"],
                                   ["Telecom noise (worst section)", fmt(max([t["ep_mV"] for t in em["telecom"]] + [0]), 3) + " mV"]]},
              {"t": "h1", "text": "Method"},
              {"t": "p", "text": f"The running rails of all tracks are taken as bonded and modelled as a ladder of rail resistance ({El['r_ret']} Ω/km per track) "
                                 f"and rail–earth conductance ({M['g_track_S_km']} S/km per track) every {M['dx']} m. Each second the simulated train currents "
                                 "are injected at the train positions and the substation return currents drawn at the substations. Touch voltage is taken "
                                 "equal to the rail potential, with the EN 50122-1 limits for the exposure time. Stray current is the average current leaving "
                                 "the rails per metre of track (EN 50122-2). The static field is from the conductor rail and return currents (Biot–Savart). "
                                 "Induced voltage in a parallel cable is 2πf·M′·L·I for the rectifier ripple harmonics, with psophometric weighting, screening "
                                 "and cable sensitivity."},
              {"t": "h1", "text": "Input data"},
              {"t": "kv", "rows": [["Rail–earth conductance", f"{M['g_track_S_km']} S/km per track (limit {M['g_limit_S_km']})"],
                                   ["Substation negative", f"earthed, {M['R_tss_earth']} Ω" if E.num(M.get("R_tss_earth"), None) is not None else "floating"],
                                   ["Conductor rail position", f"{M['cond_y']} m from track centre, {M['cond_z']} m high"],
                                   ["Field points", f"{M['field_d']} m at {M['field_h']} m height"],
                                   ["Telecom cable", f"{M['cable_d']} m from track centre, level {M['cable_z']} m, screening {M['screening']}, sensitivity {M['sensitivity']}"]]},
              {"t": "table", "head": ["Exposure s"] + [fmt(p[0], 2).rstrip("0").rstrip(".") for p in M["touch_limits"]], "rows": [["Limit V"] + [f0(p[1]) for p in M["touch_limits"]]]},
              {"t": "h1", "text": "Rail potential and touch voltage"}]
        plt, fig, ax = _fig(3.8)
        ax.plot(em["x"], em["Umax"], color=PC["acc"], lw=1.4, label="Highest")
        ax.plot(em["x"], em["Umin"], color=PC["up"], lw=1.4, label="Lowest")
        ax.plot(em["x"], em["U300"], color=PC["dn"], lw=1.2, ls="--", label="Highest 300 s average |U|")
        if Lm["s1"] is not None:
            ax.axhline(Lm["s1"], color=PC["fail"], ls="--", lw=1, label="Limit ≤ 1 s")
            ax.axhline(-Lm["s1"], color=PC["fail"], ls="--", lw=1)
        _vlines(ax, cfg); ax.set_xlabel("Chainage (m)"); ax.set_ylabel("Rail potential (V)"); ax.legend(fontsize=7, ncol=4)
        B.append(_png(plt, fig, "Rail potential envelope along the line"))
        pf = lambda v, l: "PASS" if l is None or v is None or v <= l else "FAIL"
        B += [{"t": "table", "head": ["Exposure", "Value V", "Limit V", "Result"],
               "rows": [["≤ 1 s (instantaneous)", f0(em["max1"]), f0(Lm["s1"]), pf(em["max1"], Lm["s1"])],
                        ["300 s average", f0(em["max300"]), f0(Lm["s300"]), pf(em["max300"], Lm["s300"])],
                        ["Long term", f0(em["maxMean"]), f0(Lm["long"]), pf(em["maxMean"], Lm["long"])],
                        [f"Fault, {cfg['sc']['t_clear_s']} s", f0(em["faultUrail"]), f0(Lm["fault"]), pf(em["faultUrail"], Lm["fault"])]]},
              {"t": "h1", "text": "Stray current"},
              {"t": "p", "text": f"Average current leaving the rails {f1(em['strayAvg_A'])} A over {f0((r['extent'][1] - r['extent'][0]) * em['nTracks'])} m of track: "
                                 f"{f2(em['stray_mA_m'])} mA/m against {M['stray_limit_mA_m']} mA/m. Keep the rail–earth conductance below {M['g_limit_S_km']} S/km per "
                                 "track by rail insulation and drainage design, and confirm by measurement."},
              {"t": "h1", "text": "Static magnetic field"},
              {"t": "table", "head": ["Distance m", "B at peak µT", "B at 1-h RMS µT", "Reference µT"],
               "rows": [[fmt(f["d"], 1), f1(f["Bpk_uT"]), f1(f["Brms_uT"]), f0(E.num(M.get("B_limit_mT"), 0.5) * 1000)] for f in em["field"]]},
              {"t": "h1", "text": "Inductive interference to telecom cables"},
              {"t": "p", "text": f"Mutual inductance between the traction loop and the cable {fmt(em['Mprime_uH_m'], 3)} µH/m, parallel length = feeding section of each substation."},
              {"t": "table", "head": ["Section", "Length m", "f Hz", "Ripple A", "emf mV", "Weight", "Noise mV"],
               "rows": [[t["tss"] if k == 0 else "", f0(t["Lpar"]) if k == 0 else "", f0(e["f"]), f2(e["I"]), f1(e["E_mV"]), f2(e["w"]),
                         fmt(t["ep_mV"], 3) if k == 0 else ""] for t in em["telecom"] for k, e in enumerate(t["E"])]},
              {"t": "h1", "text": "Compliance checks"}, _check_table(r, groups)]
    B.append({"t": "h1", "text": "Assumptions and limitations"})
    for p in ["Results are planning-level values from the inputs listed; confirm with the supplier's final data and site measurement.",
              "Trains are not slowed by low line voltage; their power is cut back but the timetable holds.",
              "Fault currents rise with a first-order time constant; rectifier overshoot, arc resistance and breaker dynamics are not modelled.",
              "Harmonics assume no background distortion and no network resonance.",
              "Touch voltage is taken equal to rail potential, without body or footwear resistance.",
              "The HTML model (TPS_Model.html) runs the same calculation on the same inputs file."]:
        B.append({"t": "bullet", "text": p})
    rp["blocks"] = B
    return rp


# ----------------------------------------------------------------------
# Word writer
# ----------------------------------------------------------------------
def write_docx(rp, path, engine_line):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    INK, INK2, MUTE = RGBColor(0x1C, 0x26, 0x31), RGBColor(0x4D, 0x5A, 0x67), RGBColor(0x6E, 0x76, 0x7E)
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), fill)
        tcPr.append(sh)

    def field(run, instr):
        for kind, text in (("begin", None), (None, instr), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)

    hp = sec.header.paragraphs[0]
    hr = hp.add_run(f"{rp['docref']} · {rp['title']}")
    hr.font.size, hr.font.color.rgb = Pt(8), MUTE
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run(engine_line + "   ·   Page ")
    fr.font.size, fr.font.color.rgb = Pt(8), MUTE
    r2 = fp.add_run()
    r2.font.size = Pt(8)
    field(r2, "PAGE")
    r3 = fp.add_run(" of ")
    r3.font.size, r3.font.color.rgb = Pt(8), MUTE
    r4 = fp.add_run()
    r4.font.size = Pt(8)
    field(r4, "NUMPAGES")

    def table(head, rows, kv=False):
        nc = max(len(x) for x in ([head] if head else []) + rows)
        wts = []
        for j in range(nc):
            m = min(len(str(head[j])), 12) if head else 4
            for x in rows:
                m = max(m, len(str(x[j])) if j < len(x) else 0)
            wts.append(min(max(m, 7), 60 if kv else 40) + 2)
        tot = sum(wts)
        widths = [Cm(17.4 * w / tot) for w in wts]
        t = doc.add_table(rows=0, cols=nc)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if head:
            cells = t.add_row().cells
            for j, h in enumerate(head):
                cells[j].text = ""
                run = cells[j].paragraphs[0].add_run(str(h))
                run.bold, run.font.size, run.font.color.rgb = True, Pt(8.5), RGBColor(0xFF, 0xFF, 0xFF)
                shade(cells[j], "1F3A4D")
        for i, x in enumerate(rows):
            cells = t.add_row().cells
            for j in range(nc):
                v = str(x[j]) if j < len(x) and x[j] is not None else ""
                cells[j].text = ""
                run = cells[j].paragraphs[0].add_run(v)
                run.font.size = Pt(8.5)
                if (kv and j == 0) or v in ("PASS", "FAIL"):
                    run.bold = True
                run.font.color.rgb = RGBColor(0xB3, 0x26, 0x1E) if v == "FAIL" else RGBColor(0x1B, 0x74, 0x42) if v == "PASS" else (INK2 if kv and j == 0 else INK)
                if not kv and i % 2:
                    shade(cells[j], "F4F5F0")
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        grid = t._tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[j].twips)))
        for j, col in enumerate(t.columns):
            col.width = widths[j]
        for row in t.rows:
            for j, c in enumerate(row.cells):
                c.width = widths[j]
        doc.add_paragraph()

    h1 = h2 = 0
    for b in rp["blocks"]:
        if b["t"] == "title":
            for _ in range(6):
                doc.add_paragraph()
            p = doc.add_paragraph()
            run = p.add_run(b["title"])
            run.bold, run.font.size, run.font.color.rgb = True, Pt(26), INK
            p = doc.add_paragraph()
            run = p.add_run(b["sub"])
            run.font.size, run.font.color.rgb = Pt(14), INK2
            doc.add_paragraph()
            table(None, b["meta"], kv=True)
            doc.add_page_break()
        elif b["t"] == "h1":
            h1 += 1
            h2 = 0
            hd = doc.add_heading(f"{h1}  {b['text']}", level=1)
            for run in hd.runs:
                run.font.color.rgb = INK
        elif b["t"] == "h2":
            h2 += 1
            hd = doc.add_heading(f"{h1}.{h2}  {b['text']}", level=2)
            for run in hd.runs:
                run.font.color.rgb = INK
        elif b["t"] == "p":
            doc.add_paragraph(b["text"])
        elif b["t"] == "bullet":
            doc.add_paragraph(b["text"], style="List Bullet")
        elif b["t"] == "kv":
            table(None, b["rows"], kv=True)
        elif b["t"] == "table":
            table(b["head"], b["rows"])
        elif b["t"] == "chart":
            doc.add_picture(io.BytesIO(b["png"]), width=Cm(17))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if b.get("caption"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run("Figure: " + b["caption"])
                run.italic, run.font.size, run.font.color.rgb = True, Pt(8.5), MUTE
    doc.save(path)
    return path


def write_all(cfg, r, flt, st, outdir, dc=None):
    os.makedirs(outdir, exist_ok=True)
    out = []
    for rid, code, _ in REPORTS:
        rp = build_report(rid, cfg, r, flt, st, dc)
        out.append(write_docx(rp, os.path.join(outdir, rp["file"] + ".docx"), f"{r['engine']} · {r['scenario']}"))
    return out
